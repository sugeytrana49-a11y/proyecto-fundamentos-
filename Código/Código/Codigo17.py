"""
Búsqueda de información en una factura — estudio con un único participante.

Flujo:
  Inicio -> (entrenamiento, solo la primera vez) -> [ciclo de cada prueba]
  Ciclo:  pantalla del grupo (ubicar al participante a la distancia asignada)
          -> instrucción (muestra el código a buscar)
          -> factura (el cronómetro arranca cuando ya está en pantalla y termina al hacer
             clic en la fila correcta; un clic en otra fila cuenta como error y no detiene)
          -> descanso (30 s entre fichas y 5 min a la mitad)
  No hay límite de tiempo: la prueba termina solo cuando se encuentra el código.

Teclas (para el grupo):
  Esc      salir de pantalla completa
  Ctrl+Q   cerrar el programa
  Ctrl+D   DESCARTAR la prueba actual (no se guarda; se pasa a otra ficha de repuesto)
  Ctrl+S   saltar el descanso

Factores asignados por el programa (marginales balanceadas y en orden aleatorio):
  número de líneas, tamaño de letra, contraste (razón calculada), distancia asignada.
Variable observada: posición del objetivo (fila_obj y dist_centro_filas, al azar).

Modos:  MODO = "piloto" (premuestra de 25)  |  MODO = "estudio" (muestra completa).

Uso:  python busqueda_factura.py        (requiere:  pip install pillow openpyxl)
"""
import math
import os
import random
import time
import tkinter as tk
from datetime import datetime
from tkinter import messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont, ImageTk

# En Windows con escala de pantalla (125 %, 150 %), evita que el sistema "estire" la ventana:
# así 1 píxel de la factura es 1 píxel real y las letras se ven nítidas.
try:
    import ctypes

    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    pass

# ======================================================================
# CONFIGURACIÓN GENERAL
# ======================================================================
MODO = "piloto"                      # "piloto" o "estudio"
PRUEBAS_POR_MODO = {"piloto": 25, "estudio": 100}   # cambie "estudio" según el tamaño de muestra calculado
TOTAL_PRUEBAS = PRUEBAS_POR_MODO[MODO]
REPUESTOS = 10                       # fichas extra por si se descarta alguna prueba
N_PRACTICA = 2                       # fichas de entrenamiento (no se guardan); 0 = sin entrenamiento
DESCANSO_ENTRE_S = 30                # descanso entre fichas
DESCANSO_LARGO_S = 300               # descanso a la mitad de las pruebas

# Distintas semillas para que la premuestra y el estudio no repitan fichas
SEMILLA = 2026 + (0 if MODO == "piloto" else 1000)

NIVELES_FILAS = [4, 8, 12]           # cantidad de información: líneas de la factura
NIVELES_FUENTE = [13, 17, 21, 25]    # tamaño de letra de la tabla (px)
DISTANCIAS_M = [0.5, 1.0, 1.5, 2.0]  # distancia asignada entre el participante y la pantalla

# Contraste: colores en escala de grises; la razón de contraste se calcula más abajo
CONFIG_CONTRASTES = {
    "alto": {"fondo": (255, 255, 255), "texto": (0, 0, 0)},
    "medio": {"fondo": (235, 235, 235), "texto": (70, 70, 70)},
    "bajo": {"fondo": (205, 205, 205), "texto": (110, 110, 110)},
}

# Geometría fija de la factura (px)
W, H = 1200, 800
X_TABLA_INI, X_TABLA_FIN = 80, W - 80
Y_FILA_0 = 253       # posición del texto de la primera fila
ALTO_FILA = 30

PANTALLA_COMPLETA = True
FONDO_VENTANA = "#1e1e1e"

# Un solo Excel con tres hojas: Resultados, Descartes y Diseño
ARCHIVO_RESULTADOS = f"resultados_factura_{MODO}.xlsx"
ARCHIVO_RESPALDO = f"resultados_factura_{MODO}_respaldo.xlsx"

# ======================================================================
# BANCO DE DATOS PARA LAS LÍNEAS DE LA FACTURA
# ======================================================================
DESCRIPCIONES_BASE = [
    "Servicio de consultoría técnica", "Mantenimiento preventivo de línea",
    "Suministro de válvulas de presión", "Kit de juntas tóricas de repuesto",
    "Lubricante sintético industrial", "Filtro de aire de alta eficiencia",
    "Panel de control digital v2", "Cable apantallado de cobre",
    "Sensor fotoeléctrico de proximidad", "Rodamiento de bolas sellado",
    "Correa de transmisión dentada", "Módulo de relés de seguridad",
    "Motor trifásico de 5 HP", "Variador de frecuencia 15 kW",
    "Manguera hidráulica reforzada", "Cilindro neumático doble efecto",
    "Tornillería de acero inoxidable", "Banda transportadora modular",
    "Reductor de velocidad helicoidal", "Interruptor de fin de carrera",
    "Contactor magnético 40 A", "Aceite hidráulico ISO 46",
    "Electrodo de soldadura 7018", "Guantes de seguridad de nitrilo",
    "Casco de seguridad clase E", "Gafas de protección antiempaño",
    "Etiquetas térmicas industriales", "Cinta de embalaje reforzada",
    "Pallet de madera tratada", "Sellador de roscas anaeróbico",
]

# ======================================================================
# FUENTES
# ======================================================================
_CACHE_FUENTES = {}


def fuente(tam):
    tam = int(tam)
    if tam in _CACHE_FUENTES:
        return _CACHE_FUENTES[tam]
    candidatas = [
        "arialbd.ttf", "Arial Bold.ttf", "C:/Windows/Fonts/arialbd.ttf",
        "/Library/Fonts/Arial Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ]
    for ruta in candidatas:
        try:
            _CACHE_FUENTES[tam] = ImageFont.truetype(ruta, tam)
            return _CACHE_FUENTES[tam]
        except OSError:
            continue
    raise RuntimeError("No se encontró una fuente sans-serif en negrita (Arial, Liberation Sans o DejaVu).")


# ======================================================================
# CONTRASTE (razón de contraste, fórmula WCAG)
# ======================================================================
def _luminancia(color):
    def lin(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4

    r, g, b = color
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def razon_contraste(fondo, texto):
    claro, oscuro = sorted((_luminancia(fondo), _luminancia(texto)), reverse=True)
    return round((claro + 0.05) / (oscuro + 0.05), 2)


# ======================================================================
# DISEÑO DE LAS FICHAS
# ======================================================================
def niveles_balanceados(niveles, total, rng):
    """Lista de longitud total + REPUESTOS: las primeras 'total' tienen cada nivel el mismo
    número de veces (o difieren en 1) en orden aleatorio; los repuestos van al azar."""
    base = list(niveles) * (total // len(niveles)) + rng.sample(list(niveles), total % len(niveles))
    rng.shuffle(base)
    repuestos = [rng.choice(list(niveles)) for _ in range(REPUESTOS)]
    return base + repuestos


def crear_documento(doc_id, num_filas, tam_fuente, nivel_contraste, distancia_m, rng):
    cfg = CONFIG_CONTRASTES[nivel_contraste]
    descripciones = rng.sample(DESCRIPCIONES_BASE, num_filas)
    codigos = rng.sample(range(1000, 10000), num_filas)     # códigos únicos
    fila_obj = rng.randint(0, num_filas - 1)

    lineas = []
    for i in range(num_filas):
        lineas.append({
            "codigo": f"REF-{codigos[i]}",
            "descripcion": descripciones[i],
            "unidades": str(rng.randint(1, 10)),
            "precio": f"{rng.randint(5, 250) * 100:,}".replace(",", " ") + " CRC",
        })
    return {
        "doc_id": doc_id,
        "num_filas": num_filas,
        "tamano_fuente_doc": tam_fuente,
        "nivel_contraste": nivel_contraste,
        "razon_contraste": razon_contraste(cfg["fondo"], cfg["texto"]),
        "color_fondo_rgb": cfg["fondo"],
        "color_texto_rgb": cfg["texto"],
        "distancia_asignada_m": distancia_m,
        "fila_obj": fila_obj,
        "dist_centro_filas": abs(fila_obj - (num_filas - 1) / 2),
        "valor_objetivo": lineas[fila_obj]["codigo"],
        "lineas_datos": lineas,
    }


def construir_secuencia():
    rng = random.Random(SEMILLA)
    n = TOTAL_PRUEBAS + REPUESTOS
    filas = niveles_balanceados(NIVELES_FILAS, TOTAL_PRUEBAS, rng)
    fuentes = niveles_balanceados(NIVELES_FUENTE, TOTAL_PRUEBAS, rng)
    contrastes = niveles_balanceados(list(CONFIG_CONTRASTES), TOTAL_PRUEBAS, rng)
    distancias = niveles_balanceados(DISTANCIAS_M, TOTAL_PRUEBAS, rng)
    return [
        crear_documento(f"FAC-{SEMILLA}/{i + 1:03d}", filas[i], fuentes[i], contrastes[i], distancias[i], rng)
        for i in range(n)
    ]


def construir_practica():
    rng = random.Random(SEMILLA + 99)
    return [
        crear_documento(f"PRACTICA-{i + 1}", rng.choice(NIVELES_FILAS), rng.choice(NIVELES_FUENTE),
                        rng.choice(list(CONFIG_CONTRASTES)), DISTANCIAS_M[i % len(DISTANCIAS_M)], rng)
        for i in range(N_PRACTICA)
    ]


def validar_textos():
    """Comprueba al iniciar que ningún texto se sale de su columna con la letra más grande."""
    medidor = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    f = fuente(max(NIVELES_FUENTE))
    anchos = {
        "descripción": (max(medidor.textlength(t, font=f) for t in DESCRIPCIONES_BASE), 700 - 260 - 10),
        "código": (medidor.textlength("REF-9999", font=f), 260 - 100 - 10),
        "precio": (medidor.textlength("25 000 CRC", font=f), X_TABLA_FIN - 850 - 10),
    }
    for nombre, (ancho, maximo) in anchos.items():
        if ancho > maximo:
            raise ValueError(f"La columna '{nombre}' se desborda con letra de {max(NIVELES_FUENTE)} px "
                             f"({ancho:.0f} px > {maximo} px). Acorte los textos.")


# ======================================================================
# DIBUJO
# ======================================================================
def dibujar_documento(d):
    img = Image.new("RGB", (W, H), d["color_fondo_rgb"])
    draw = ImageDraw.Draw(img)
    col = d["color_texto_rgb"]
    f_titulo, f_sub, f_bold = fuente(28), fuente(13), fuente(11)

    draw.text((80, 40), "FACTURA", fill=col, font=f_titulo)
    draw.text((80, 80), "Fecha: 10/11/2026", fill=col, font=f_sub)
    draw.text((80, 100), f"Nº factura: {d['doc_id']}", fill=col, font=f_sub)

    draw.text((80, 140), "DATOS EMISOR", fill=col, font=f_bold)
    draw.text((80, 160), "• Industrial S.A. - Céd. jurídica: 3-101-000000", fill=col, font=f_sub)
    draw.text((80, 178), "• Parque Industrial, bodega 25", fill=col, font=f_sub)

    draw.text((500, 140), "DATOS CLIENTE", fill=col, font=f_bold)
    draw.text((500, 160), "• Cliente Comercial S.A.", fill=col, font=f_sub)
    draw.text((500, 178), "• Dirección de entrega sector B", fill=col, font=f_sub)

    y_tabla = 220
    draw.line([(X_TABLA_INI, y_tabla), (X_TABLA_FIN, y_tabla)], fill=(150, 150, 150), width=2)
    for x, titulo in ((100, "Código"), (260, "Descripción"), (700, "Unidades"), (850, "Precio")):
        draw.text((x, y_tabla + 8), titulo, fill=col, font=f_bold)
    draw.line([(X_TABLA_INI, y_tabla + 30), (X_TABLA_FIN, y_tabla + 30)], fill=(180, 180, 180), width=1)

    f_celda = fuente(d["tamano_fuente_doc"])
    y = Y_FILA_0
    for linea in d["lineas_datos"]:
        draw.text((100, y), linea["codigo"], fill=col, font=f_celda)
        draw.text((260, y), linea["descripcion"], fill=col, font=f_celda)
        draw.text((700, y), linea["unidades"], fill=col, font=f_celda)
        draw.text((850, y), linea["precio"], fill=col, font=f_celda)
        y += ALTO_FILA
        draw.line([(X_TABLA_INI, y - 2), (X_TABLA_FIN, y - 2)], fill=(210, 210, 210), width=1)
    return img


# ======================================================================
# ARCHIVO EXCEL (cada columna lleva su categoría en la fila superior)
# ======================================================================
HOJA_RESULTADOS = "Resultados"
HOJA_DESCARTES = "Descartes"
HOJA_DISENO = "Diseño"

COLOR_CATEGORIA = {
    "Identificación": "D9D9D9",
    "Factor asignado": "BDD7EE",
    "Variable observada": "C6E0B4",
    "Respuesta medida": "FFE699",
    "Control del procedimiento": "E4DFEC",
}

# (categoría, título de la columna, clave del dato)
ESTRUCTURA = {
    HOJA_RESULTADOS: [
        ("Identificación", "ID prueba", "id_prueba"),
        ("Identificación", "ID factura", "doc_id"),
        ("Identificación", "Bloque", "bloque"),
        ("Factor asignado", "Número de líneas", "num_filas"),
        ("Factor asignado", "Tamaño de letra (px)", "tamano_fuente_doc"),
        ("Factor asignado", "Nivel de contraste", "nivel_contraste"),
        ("Factor asignado", "Razón de contraste", "razon_contraste"),
        ("Factor asignado", "Distancia asignada (m)", "distancia_asignada_m"),
        ("Variable observada", "Código objetivo", "valor_objetivo"),
        ("Variable observada", "Fila del objetivo", "fila_obj"),
        ("Variable observada", "Distancia al centro (filas)", "dist_centro_filas"),
        ("Respuesta medida", "Tiempo (s)", "tiempo_s"),
        ("Respuesta medida", "Errores (n)", "errores"),
        ("Control del procedimiento", "Distancia real (m)", "distancia_real_m"),
        ("Control del procedimiento", "Pantalla (px)", "pantalla_px"),
        ("Control del procedimiento", "Fecha y hora", "hora"),
    ],
    HOJA_DESCARTES: [
        ("Identificación", "ID prueba", "id_prueba"),
        ("Identificación", "ID factura", "doc_id"),
        ("Respuesta medida", "Errores (n)", "errores"),
        ("Control del procedimiento", "Fecha y hora", "hora"),
    ],
    HOJA_DISENO: [
        ("Identificación", "Orden", "orden"),
        ("Identificación", "ID factura", "doc_id"),
        ("Identificación", "Repuesto (1 = sí)", "repuesto"),
        ("Factor asignado", "Número de líneas", "num_filas"),
        ("Factor asignado", "Tamaño de letra (px)", "tamano_fuente_doc"),
        ("Factor asignado", "Nivel de contraste", "nivel_contraste"),
        ("Factor asignado", "Razón de contraste", "razon_contraste"),
        ("Factor asignado", "Distancia asignada (m)", "distancia_asignada_m"),
        ("Variable observada", "Código objetivo", "valor_objetivo"),
        ("Variable observada", "Fila del objetivo", "fila_obj"),
        ("Variable observada", "Distancia al centro (filas)", "dist_centro_filas"),
    ],
}


def _crear_hoja(wb, nombre):
    ws = wb.create_sheet(nombre)
    est = ESTRUCTURA[nombre]
    lado = Side(style="thin", color="999999")
    for i, (cat, titulo, clave) in enumerate(est, start=1):
        for fila_enc, texto in ((1, cat), (2, titulo)):
            c = ws.cell(row=fila_enc, column=i, value=texto)
            c.fill = PatternFill("solid", fgColor=COLOR_CATEGORIA[cat])
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=lado, right=lado, top=lado, bottom=lado)
        ws.column_dimensions[get_column_letter(i)].width = 22 if clave == "hora" else max(len(titulo) + 4, 16)
    i = 0
    while i < len(est):          # une las celdas de categoría consecutivas iguales
        j = i
        while j + 1 < len(est) and est[j + 1][0] == est[i][0]:
            j += 1
        if j > i:
            ws.merge_cells(start_row=1, start_column=i + 1, end_row=1, end_column=j + 1)
        i = j + 1
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"
    return ws


def _abrir_libro(archivo):
    existe = os.path.exists(archivo)
    wb = load_workbook(archivo) if existe else Workbook()
    if not existe:
        wb.remove(wb.active)
    for nombre in ESTRUCTURA:
        if nombre not in wb.sheetnames:
            _crear_hoja(wb, nombre)
    return wb


def _guardar_libro(wb, archivo):
    tmp = archivo + ".tmp"      # se guarda aparte y luego se reemplaza, para no dañar el archivo
    wb.save(tmp)
    os.replace(tmp, archivo)


def contar_filas(hoja, archivo=None):
    archivo = archivo or ARCHIVO_RESULTADOS
    if not os.path.exists(archivo):
        return 0
    wb = load_workbook(archivo)
    try:
        return max(wb[hoja].max_row - 2, 0) if hoja in wb.sheetnames else 0
    finally:
        wb.close()


def agregar_fila(hoja, fila, archivo=None):
    archivo = archivo or ARCHIVO_RESULTADOS
    wb = _abrir_libro(archivo)
    wb[hoja].append([fila[clave] for _, _, clave in ESTRUCTURA[hoja]])
    _guardar_libro(wb, archivo)


def escribir_diseno(secuencia):
    wb = _abrir_libro(ARCHIVO_RESULTADOS)
    ws = wb[HOJA_DISENO]
    if ws.max_row > 2:          # el diseño ya se escribió antes
        return
    for i, d in enumerate(secuencia, start=1):
        fila = {
            "orden": i, "doc_id": d["doc_id"], "repuesto": int(i > TOTAL_PRUEBAS),
            "num_filas": d["num_filas"], "tamano_fuente_doc": d["tamano_fuente_doc"],
            "nivel_contraste": d["nivel_contraste"], "razon_contraste": d["razon_contraste"],
            "distancia_asignada_m": d["distancia_asignada_m"], "fila_obj": d["fila_obj"] + 1,
            "dist_centro_filas": d["dist_centro_filas"], "valor_objetivo": d["valor_objetivo"],
        }
        ws.append([fila[clave] for _, _, clave in ESTRUCTURA[HOJA_DISENO]])
    _guardar_libro(wb, ARCHIVO_RESULTADOS)


# ======================================================================
# APLICACIÓN
# ======================================================================
class App:
    def __init__(self, secuencia, practica):
        self.secuencia = secuencia
        self.docs_practica = practica

        self.root = tk.Tk()
        self.root.title("Estudio de búsqueda de información en factura")
        self.root.configure(bg=FONDO_VENTANA)
        if PANTALLA_COMPLETA:
            try:
                self.root.state("zoomed")  # solo Windows
            except tk.TclError:
                pass
            self.root.attributes("-fullscreen", True)
        else:
            self.root.geometry("1400x900")
        self.root.bind("<Escape>", lambda e: self.root.attributes("-fullscreen", False))
        self.root.bind("<Control-q>", lambda e: self.cerrar())
        self.root.protocol("WM_DELETE_WINDOW", self.cerrar)
        self.root.bind("<Control-d>", lambda e: self.descartar())
        self.root.bind("<Control-s>", lambda e: self.saltar_descanso())

        self.panel = None
        self.timer_id = None
        self.pendientes = set()
        self.estado = None
        self.activo = False
        self.practica = False
        self.idx_practica = 0
        self.errores = 0
        self.foto = None

        self.guardadas = self.con_reintento(lambda: contar_filas(HOJA_RESULTADOS), "leer los resultados")
        self.descartadas = self.con_reintento(lambda: contar_filas(HOJA_DESCARTES), "leer los descartes")
        self.pos = self.guardadas + self.descartadas
        self.con_reintento(lambda: escribir_diseno(self.secuencia), "guardar el diseño")

        self.root.update_idletasks()
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        if PANTALLA_COMPLETA and (sw < W or sh < H):
            self.pantalla_insuficiente(sw, sh)
        else:
            self.pantalla_inicio()

    # ---------- utilidades ----------
    def con_reintento(self, accion, descripcion):
        while True:
            try:
                return accion()
            except PermissionError:
                if not messagebox.askretrycancel(
                        "Archivo en uso",
                        f"No se pudo {descripcion}.\nCierre el archivo si lo tiene abierto en Excel y presione Reintentar.",
                        parent=self.root):
                    raise

    def despues(self, ms, funcion):
        """Como root.after, pero se puede cancelar al cerrar (evita avisos de error en consola)."""
        def ejecutar():
            self.pendientes.discard(id_)
            funcion()
        id_ = self.root.after(ms, ejecutar)
        self.pendientes.add(id_)

    def cerrar(self):
        for id_ in list(self.pendientes):
            self.root.after_cancel(id_)
        if self.timer_id is not None:
            self.root.after_cancel(self.timer_id)
        self.root.destroy()

    def limpiar(self):
        if self.timer_id is not None:
            self.root.after_cancel(self.timer_id)
            self.timer_id = None
        if self.panel is not None:
            self.panel.destroy()
        self.panel = tk.Frame(self.root, bg=FONDO_VENTANA)
        self.panel.pack(fill="both", expand=True)
        self.root.unbind("<space>")
        self.estado = None
        self.activo = False
        return self.panel

    def etiqueta(self, parent, texto, tam=16, color="white", negrita=False):
        return tk.Label(parent, text=texto, bg=FONDO_VENTANA, fg=color,
                        font=("Arial", tam, "bold" if negrita else "normal"))

    def centro(self):
        p = self.limpiar()
        c = tk.Frame(p, bg=FONDO_VENTANA)
        c.place(relx=0.5, rely=0.5, anchor="center")
        return c

    # ---------- pantallas ----------
    def pantalla_insuficiente(self, ancho, alto):
        c = self.centro()
        self.etiqueta(c, "La pantalla es demasiado pequeña", 28, "#ff7b7b", True).pack(pady=14)
        self.etiqueta(c, f"Disponible: {ancho} × {alto} px. La factura necesita {W} × {H} px.", 16).pack(pady=6)
        self.etiqueta(c, "Use un monitor más grande o baje la escala de pantalla de Windows al 100 %.",
                      14, "#aaaaaa").pack(pady=6)
        tk.Button(c, text="Cerrar", font=("Arial", 16, "bold"), command=self.cerrar).pack(pady=20)

    def pantalla_inicio(self):
        if self.guardadas >= TOTAL_PRUEBAS:
            return self.pantalla_fin_total()
        c = self.centro()
        self.etiqueta(c, "Búsqueda de información en una factura", 30, "#7fb2ff", True).pack(pady=14)
        self.etiqueta(c, f"Modo: {MODO}  |  Pruebas guardadas: {self.guardadas} de {TOTAL_PRUEBAS}", 16).pack(pady=6)
        self.etiqueta(c, "En cada prueba verá una factura. Encuentre el código que se le indique y haga clic en su fila.",
                      15, "#cccccc").pack(pady=6)
        entrena = self.guardadas == 0 and N_PRACTICA > 0
        texto = f"Comenzar con {N_PRACTICA} fichas de práctica" if entrena else "Continuar"
        tk.Button(c, text=texto, font=("Arial", 18, "bold"),
                  command=self.iniciar_practica if entrena else self.pantalla_grupo).pack(pady=26)

    # --- entrenamiento ---
    def iniciar_practica(self):
        self.practica = True
        self.idx_practica = 0
        self.siguiente_practica()

    def siguiente_practica(self):
        if self.idx_practica >= len(self.docs_practica):
            return self.pantalla_fin_practica()
        self.pantalla_grupo()

    def tras_practica(self):
        if self.idx_practica >= len(self.docs_practica):
            return self.pantalla_fin_practica()
        self.descanso(DESCANSO_ENTRE_S, False)

    def pantalla_fin_practica(self):
        self.practica = False
        c = self.centro()
        self.etiqueta(c, "Entrenamiento completo", 32, "#7fe08a", True).pack(pady=16)
        self.etiqueta(c, "Ahora empieza la prueba real. A partir de aquí los tiempos se guardan.", 16).pack(pady=8)
        tk.Button(c, text="Iniciar prueba real", font=("Arial", 18, "bold"), command=self.pantalla_grupo).pack(pady=24)

    # --- pantalla del grupo ---
    def pantalla_grupo(self):
        if self.practica:
            self.diseno = self.docs_practica[self.idx_practica]
            titulo = f"Práctica {self.idx_practica + 1} de {len(self.docs_practica)}"
        else:
            if self.guardadas >= TOTAL_PRUEBAS:
                return self.pantalla_fin_total()
            if self.pos >= len(self.secuencia):
                c = self.centro()
                self.etiqueta(c, "Se acabaron las fichas de repuesto", 24, "#ff7b7b", True).pack(pady=14)
                tk.Button(c, text="Cerrar", font=("Arial", 16, "bold"), command=self.cerrar).pack(pady=14)
                return
            self.diseno = self.secuencia[self.pos]
            self.id_prueba = self.guardadas + 1
            titulo = f"Prueba {self.id_prueba} de {TOTAL_PRUEBAS}"
        c = self.centro()
        self.etiqueta(c, titulo, 24, "#7fb2ff", True).pack(pady=(0, 8))
        self.etiqueta(c, "(Pantalla para el grupo)", 12, "#aaaaaa").pack(pady=(0, 18))
        self.etiqueta(c, f"Ubique al participante a {self.diseno['distancia_asignada_m']:.1f} m de la pantalla",
                      22).pack(pady=10)
        fila = tk.Frame(c, bg=FONDO_VENTANA)
        fila.pack(pady=14)
        self.etiqueta(fila, "Distancia real (m):", 14).grid(row=0, column=0, padx=10)
        self.e_dist = tk.Entry(fila, font=("Arial", 14), width=8)
        self.e_dist.insert(0, f"{self.diseno['distancia_asignada_m']:.1f}")
        self.e_dist.grid(row=0, column=1)
        self.lbl_error = self.etiqueta(c, "", 12, "#ff7b7b")
        self.lbl_error.pack()
        tk.Button(c, text="Continuar", font=("Arial", 14, "bold"), command=self.confirmar_distancia).pack(pady=14)
        self.e_dist.focus_set()

    def confirmar_distancia(self):
        try:
            dist = float(self.e_dist.get().strip().replace(",", "."))
            if not 0.2 <= dist <= 5.0:
                raise ValueError("fuera de rango (0,2 a 5 m)")
        except ValueError as err:
            self.lbl_error.config(text=f"Distancia no válida: {err}")
            return
        self.distancia_real = dist
        self.mostrar_instruccion()

    # --- instrucción y factura ---
    def mostrar_instruccion(self):
        c = self.centro()
        self.estado = "instruccion"
        titulo = (f"Práctica {self.idx_practica + 1} de {len(self.docs_practica)}" if self.practica
                  else f"Prueba {self.id_prueba} de {TOTAL_PRUEBAS}")
        self.etiqueta(c, titulo, 14, "#aaaaaa").pack(pady=(0, 10))
        self.etiqueta(c, "En la factura, busque el siguiente código:", 22).pack(pady=10)
        self.etiqueta(c, self.diseno["valor_objetivo"], 54, "#ffd34d", True).pack(pady=16)
        self.etiqueta(c, "Haga clic sobre su fila lo más rápido posible.", 18, "#cccccc").pack(pady=10)
        tk.Button(c, text="Comenzar (ESPACIO)", font=("Arial", 18, "bold"),
                  command=self.mostrar_documento).pack(pady=28)
        self.root.bind("<space>", lambda e: self.mostrar_documento())

    def mostrar_documento(self):
        p = self.limpiar()
        canvas = tk.Canvas(p, bg=FONDO_VENTANA, highlightthickness=0, cursor="hand2")
        canvas.pack(fill="both", expand=True)
        self.root.update_idletasks()
        cw, ch = canvas.winfo_width(), canvas.winfo_height()
        if cw < W or ch < H:
            return self.pantalla_insuficiente(cw, ch)
        self.canvas = canvas
        self.pantalla = f"{cw}x{ch}"
        self.off_x, self.off_y = (cw - W) // 2, (ch - H) // 2
        self.foto = ImageTk.PhotoImage(dibujar_documento(self.diseno))
        canvas.create_image(self.off_x, self.off_y, anchor="nw", image=self.foto)
        canvas.bind("<Button-1>", self.clic)
        self.errores = 0
        self.estado = "documento"
        self.activo = True
        self.root.update()                  # la factura ya está pintada...
        self.t0 = time.perf_counter()       # ...y recién ahora arranca el cronómetro

    def fila_en(self, event):
        x, y = event.x - self.off_x, event.y - self.off_y
        if not X_TABLA_INI <= x <= X_TABLA_FIN:
            return None
        fila = math.floor((y - (Y_FILA_0 - 2)) / ALTO_FILA)
        return fila if 0 <= fila < len(self.diseno["lineas_datos"]) else None

    def marcar_fila(self, fila, color, ms):
        x1 = self.off_x + X_TABLA_INI
        y1 = self.off_y + (Y_FILA_0 - 2) + fila * ALTO_FILA
        marca = self.canvas.create_rectangle(x1, y1, self.off_x + X_TABLA_FIN, y1 + ALTO_FILA,
                                             outline=color, width=3)
        canvas = self.canvas

        def borrar():
            try:
                canvas.delete(marca)
            except tk.TclError:
                pass

        self.despues(ms, borrar)

    def clic(self, event):
        if not self.activo:
            return
        t = time.perf_counter() - self.t0
        fila = self.fila_en(event)
        if fila is None:
            return
        if fila == self.diseno["fila_obj"]:
            self.activo = False
            self.marcar_fila(fila, "#43a047", 600)
            if self.practica:
                self.idx_practica += 1
                self.despues(700, self.tras_practica)
            else:
                self.guardar(t)
                self.despues(400, self.tras_prueba)
        else:
            self.errores += 1
            self.marcar_fila(fila, "#e53935", 350)

    # ---------- guardado ----------
    def guardar(self, tiempo):
        d = self.diseno
        fila = {
            "id_prueba": self.id_prueba,
            "doc_id": d["doc_id"],
            "bloque": 1 if self.id_prueba <= TOTAL_PRUEBAS // 2 else 2,
            "num_filas": d["num_filas"],
            "fila_obj": d["fila_obj"] + 1,
            "dist_centro_filas": d["dist_centro_filas"],
            "valor_objetivo": d["valor_objetivo"],
            "nivel_contraste": d["nivel_contraste"],
            "razon_contraste": d["razon_contraste"],
            "tamano_fuente_doc": d["tamano_fuente_doc"],
            "distancia_asignada_m": d["distancia_asignada_m"],
            "distancia_real_m": self.distancia_real,
            "tiempo_s": round(tiempo, 3),
            "errores": self.errores,
            "hora": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "pantalla_px": self.pantalla,
        }
        while True:
            try:
                agregar_fila(HOJA_RESULTADOS, fila)
                break
            except PermissionError:
                if not messagebox.askretrycancel(
                        "Archivo en uso",
                        "No se pudo guardar el resultado.\nCierre el archivo si lo tiene abierto en Excel y presione Reintentar.\n"
                        "(Si cancela, se guarda en un archivo de respaldo.)", parent=self.root):
                    agregar_fila(HOJA_RESULTADOS, fila, ARCHIVO_RESPALDO)
                    break
        self.guardadas += 1
        self.pos += 1

    def descartar(self):
        if self.practica or self.estado not in ("instruccion", "documento"):
            return
        if not messagebox.askyesno("Descartar prueba",
                                   "¿Descartar esta prueba (no se guarda) y pasar a otra ficha?", parent=self.root):
            return
        fila = {"id_prueba": self.id_prueba, "doc_id": self.diseno["doc_id"], "errores": self.errores,
                "hora": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
        self.con_reintento(lambda: agregar_fila(HOJA_DESCARTES, fila), "guardar el descarte")
        self.activo = False
        self.descartadas += 1
        self.pos += 1
        self.pantalla_grupo()

    # ---------- descansos ----------
    def tras_prueba(self):
        if self.guardadas >= TOTAL_PRUEBAS:
            return self.pantalla_fin_total()
        largo = self.guardadas == TOTAL_PRUEBAS // 2
        self.descanso(DESCANSO_LARGO_S if largo else DESCANSO_ENTRE_S, largo)

    def descanso(self, segundos, largo):
        c = self.centro()
        self.estado = "descanso"
        self.etiqueta(c, "Descanso largo" if largo else "Descanso", 30, "#7fb2ff", True).pack(pady=10)
        self.lbl_cuenta = self.etiqueta(c, "", 72, "#ffd34d", True)
        self.lbl_cuenta.pack(pady=14)
        fichas, i = (self.docs_practica, self.idx_practica) if self.practica else (self.secuencia, self.pos)
        if i < len(fichas):
            prox = fichas[i]["distancia_asignada_m"]
            self.etiqueta(c, f"Próxima distancia: {prox:.1f} m (el grupo reubica al participante)", 14, "#aaaaaa").pack(pady=6)
        self.etiqueta(c, "Ctrl+S: saltar el descanso (solo el grupo)", 11, "#777777").pack(pady=(20, 0))
        self.restante = segundos
        self.tic()

    def tic(self):
        if self.restante <= 0:
            return self.pantalla_grupo()
        m, s = divmod(self.restante, 60)
        self.lbl_cuenta.config(text=f"{m}:{s:02d}")
        self.restante -= 1
        self.timer_id = self.root.after(1000, self.tic)

    def saltar_descanso(self):
        if self.estado == "descanso":
            self.pantalla_grupo()

    # ---------- final ----------
    def pantalla_fin_total(self):
        c = self.centro()
        self.etiqueta(c, "¡Estudio finalizado!", 36, "#7fe08a", True).pack(pady=18)
        self.etiqueta(c, f"Los datos están en {ARCHIVO_RESULTADOS}", 16).pack(pady=8)
        tk.Button(c, text="Cerrar", font=("Arial", 16, "bold"), command=self.cerrar).pack(pady=22)

    def ejecutar(self):
        self.root.mainloop()


def main():
    validar_textos()
    App(construir_secuencia(), construir_practica()).ejecutar()


if __name__ == "__main__":
    main()